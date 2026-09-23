#!/usr/bin/env python3
"""Generate the academic Hebrew DOCX for The Bridge from FINAL_REPORT_HE.md."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Emu, Mm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
MD_PATH = ROOT / "docs" / "FINAL_REPORT_HE.md"
OUT_PATH = ROOT / "docs" / "The_Bridge_Final_Report_Corrected.docx"
RENDERED = ROOT / "docs" / "diagrams" / "rendered"

HE_FONT = "David"
EN_FONT = "Arial"
CODE_FONT = "Courier New"

NSMAP_XML = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

FIGURE_META = {
    "01-architecture": (
        "6.1",
        "ארכיטקטורת הסוכנים, ה-Adapter, ה-Bridge ועמיתי ההדגמה",
        True,
    ),
    "02-sequence": (
        "6.2",
        "תרשים רצף: שליחת פעולה, Workflow והמשך הסוכן",
        True,
    ),
    "05-components": (
        "7.1",
        "דיאגרמת רכיבים ומחלקות",
        True,
    ),
    "03-workflow-states": (
        "8.1",
        "מצבי Workflow בזרימה החיה",
        False,
    ),
    "04-erd": (
        "10.1",
        "דיאגרמת ERD של טבלאות actions ו-workflows",
        False,
    ),
}

FRONT_HEADINGS = {"תקציר", "Abstract", "תודות"}
TOC_HEADINGS = {"תוכן עניינים", "רשימת איורים", "רשימת טבלאות"}
SKIP_BODY_RE = re.compile(
    r"^\*?יווצר אוטומטית|^\(יווצר אוטומטית",
)

INLINE_RE = re.compile(
    r"(https?://[^\s]+)|(`[^`]+`)|(\*\*[^*]+?\*\*)|(\*[^*]+?\*)|"
    r"(\[[^\]]+\]\([^\)]+\))|([A-Za-z][A-Za-z0-9_./:%+\-@]*)|([^`*A-Za-z]+)",
)


def qn_w(tag: str) -> str:
    return qn(f"w:{tag}")


def set_element_attr(el, name: str, value: str) -> None:
    el.set(qn(f"w:{name}"), value)


def ensure_ppr(paragraph: Paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    return pPr


def set_bidi(paragraph: Paragraph, rtl: bool) -> None:
    pPr = ensure_ppr(paragraph)
    for child in list(pPr):
        if child.tag == qn_w("bidi"):
            pPr.remove(child)
    bidi = OxmlElement("w:bidi")
    set_element_attr(bidi, "val", "1" if rtl else "0")
    pPr.append(bidi)
    if paragraph.alignment is None:
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
        )


def set_run_rtl(run, rtl: bool) -> None:
    rPr = run._r.get_or_add_rPr()
    for child in list(rPr):
        if child.tag == qn_w("rtl"):
            rPr.remove(child)
    rtl_el = OxmlElement("w:rtl")
    set_element_attr(rtl_el, "val", "1" if rtl else "0")
    rPr.append(rtl_el)
    cs = OxmlElement("w:cs")
    rPr.append(cs)


def set_run_fonts(run, *, ascii_font=EN_FONT, cs_font=HE_FONT, size_pt=12, bold=None, italic=None):
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rPr = run._r.get_or_add_rPr()
    for child in list(rPr):
        if child.tag == qn_w("rFonts"):
            rPr.remove(child)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:cs"), cs_font)
    rFonts.set(qn("w:eastAsia"), ascii_font)
    rPr.insert(0, rFonts)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "en-US")
    lang.set(qn("w:bidi"), "he-IL")
    rPr.append(lang)
    sz = rPr.find(qn_w("sz"))
    if sz is None:
        sz = OxmlElement("w:sz")
        rPr.append(sz)
    set_element_attr(sz, "val", str(int(size_pt * 2)))
    szCs = rPr.find(qn_w("szCs"))
    if szCs is None:
        szCs = OxmlElement("w:szCs")
        rPr.append(szCs)
    set_element_attr(szCs, "val", str(int(size_pt * 2)))


def apply_style_fonts(style, *, size_pt, bold, cs_font=HE_FONT, ascii_font=EN_FONT, rtl=True):
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn_w("rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:cs"), cs_font)
    pPr = style.element.get_or_add_pPr()
    bidi = pPr.find(qn_w("bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    set_element_attr(bidi, "val", "1" if rtl else "0")
    jc = pPr.find(qn_w("jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    set_element_attr(jc, "val", "right" if rtl else "left")
    spacing = pPr.find(qn_w("spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    # 1.5 lines ≈ 360 twips if 240 is single
    spacing.set(qn("w:line"), "360")
    spacing.set(qn("w:lineRule"), "auto")


def configure_section(section, *, landscape=False):
    section.page_width = Mm(297) if landscape else Mm(210)
    section.page_height = Mm(210) if landscape else Mm(297)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.top_margin = Cm(3.18)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    sectPr = section._sectPr
    pgSz = sectPr.find(qn_w("pgSz"))
    if pgSz is None:
        pgSz = OxmlElement("w:pgSz")
        sectPr.append(pgSz)
    pgSz.set(qn("w:w"), str(int(section.page_width.twips)))
    pgSz.set(qn("w:h"), str(int(section.page_height.twips)))
    pgSz.set(qn("w:orient"), "landscape" if landscape else "portrait")
    bidi = sectPr.find(qn_w("bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        sectPr.append(bidi)
    set_element_attr(bidi, "val", "1")
    docGrid = sectPr.find(qn_w("docGrid"))
    if docGrid is None:
        docGrid = OxmlElement("w:docGrid")
        sectPr.append(docGrid)
        set_element_attr(docGrid, "type", "lines")
        set_element_attr(docGrid, "linePitch", "360")


def add_field(paragraph: Paragraph, instruction: str, placeholder: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    set_element_attr(fld_begin, "fldCharType", "begin")
    run._r.append(fld_begin)

    run2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run2._r.append(instr)
    set_run_fonts(run2, size_pt=12)

    run3 = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    set_element_attr(fld_sep, "fldCharType", "separate")
    run3._r.append(fld_sep)

    run4 = paragraph.add_run(placeholder)
    set_run_fonts(run4, size_pt=12, italic=True)

    run5 = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    set_element_attr(fld_end, "fldCharType", "end")
    run5._r.append(fld_end)


def add_page_number(paragraph: Paragraph) -> None:
    add_field(paragraph, " PAGE ", "")


def set_update_fields_on_open(document: Document) -> None:
    settings = document.settings.element
    update = settings.find(qn_w("updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    set_element_attr(update, "val", "true")


def setup_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    apply_style_fonts(normal, size_pt=12, bold=False, rtl=True)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for name, size, before, after in (
        ("Heading 1", 16, 18, 12),
        ("Heading 2", 14, 14, 8),
        ("Heading 3", 12, 12, 6),
    ):
        style = styles[name]
        apply_style_fonts(style, size_pt=size, bold=True, rtl=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.keep_with_next = True
        style.font.bold = True

    def add_par_style(name, base, size, bold=False, rtl=True, after=8):
        if name in [s.name for s in styles]:
            style = styles[name]
        else:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = styles[base]
        apply_style_fonts(style, size_pt=size, bold=bold, rtl=rtl)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.5
        return style

    add_par_style("CoverTitle", "Normal", 22, bold=True, rtl=False, after=12)
    styles["CoverTitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_par_style("CoverLine", "Normal", 14, bold=False, rtl=True, after=6)
    styles["CoverLine"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_par_style("EnglishBody", "Normal", 12, rtl=False, after=8)
    styles["EnglishBody"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_par_style("Table Caption", "Normal", 11, bold=True, rtl=True, after=6)
    add_par_style("Figure Caption", "Normal", 11, bold=True, rtl=True, after=12)
    add_par_style("TOCHeading", "Heading 1", 16, bold=True, rtl=True, after=12)
    # Keep TOC headings out of the generated TOC outline
    toc_style = styles["TOCHeading"]
    pPr = toc_style.element.get_or_add_pPr()
    outline = pPr.find(qn_w("outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        pPr.append(outline)
    set_element_attr(outline, "val", "9")
    add_par_style("CodeBlock", "Normal", 10, rtl=False, after=10)
    styles["CodeBlock"].paragraph_format.line_spacing = 1.0
    styles["CodeBlock"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_par_style("ListBulletHe", "Normal", 12, rtl=True, after=4)
    add_par_style("Note", "Normal", 10, rtl=True, after=8)


def shade_cell(cell, color_hex: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    set_element_attr(shd, "fill", color_hex)
    set_element_attr(shd, "val", "clear")
    tcPr.append(shd)


def set_table_bidi(table: Table) -> None:
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    bidi = OxmlElement("w:bidiVisual")
    set_element_attr(bidi, "val", "1")
    tblPr.append(bidi)


def prevent_table_split(table: Table) -> None:
    for row in table.rows:
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)


def add_mixed_text(paragraph: Paragraph, text: str, *, size_pt=12, rtl_paragraph=True, base_bold=False):
    if not text:
        return
    for match in INLINE_RE.finditer(text):
        url, code, bold, italic, mdlink, english, rest = match.groups()
        if url:
            run = paragraph.add_run(url)
            set_run_fonts(run, ascii_font=EN_FONT, cs_font=EN_FONT, size_pt=size_pt, bold=base_bold)
            set_run_rtl(run, False)
            run.font.color.rgb = RGBColor(0x05, 0x4A, 0x91)
            continue
        if code:
            inner = code.strip("`")
            run = paragraph.add_run(inner)
            set_run_fonts(run, ascii_font=CODE_FONT, cs_font=CODE_FONT, size_pt=size_pt - 1)
            set_run_rtl(run, False)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue
        if bold:
            inner = bold[2:-2]
            add_mixed_text(paragraph, inner, size_pt=size_pt, rtl_paragraph=rtl_paragraph, base_bold=True)
            continue
        if italic:
            inner = italic[1:-1]
            run = paragraph.add_run(inner)
            has_he = bool(re.search(r"[\u0590-\u05FF]", inner))
            set_run_fonts(
                run,
                ascii_font=EN_FONT,
                cs_font=HE_FONT if has_he else EN_FONT,
                size_pt=size_pt,
                bold=base_bold,
                italic=True,
            )
            set_run_rtl(run, has_he)
            continue
        if mdlink:
            run = paragraph.add_run(mdlink)
            set_run_fonts(run, size_pt=size_pt, bold=base_bold)
            set_run_rtl(run, False)
            continue
        if english:
            run = paragraph.add_run(english)
            set_run_fonts(run, ascii_font=EN_FONT, cs_font=EN_FONT, size_pt=size_pt, bold=base_bold)
            set_run_rtl(run, False)
            continue
        if rest:
            run = paragraph.add_run(rest)
            has_he = bool(re.search(r"[\u0590-\u05FF]", rest))
            set_run_fonts(
                run,
                ascii_font=EN_FONT,
                cs_font=HE_FONT if has_he else EN_FONT,
                size_pt=size_pt,
                bold=base_bold,
            )
            set_run_rtl(run, has_he if rtl_paragraph else False)


def add_body_paragraph(document: Document, text: str, *, style="Normal", rtl=True, size_pt=12):
    p = document.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    set_bidi(p, rtl)
    p.paragraph_format.line_spacing = 1.5
    add_mixed_text(p, text, size_pt=size_pt, rtl_paragraph=rtl)
    return p


def heading_paragraph(document: Document, text: str, level: int, *, page_break=False):
    style = f"Heading {level}"
    p = document.add_paragraph(style=style)
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_bidi(p, True)
    if page_break:
        p.paragraph_format.page_break_before = True
    size = {1: 16, 2: 14, 3: 12}[level]
    add_mixed_text(p, text, size_pt=size, rtl_paragraph=True, base_bold=True)
    for run in p.runs:
        run.bold = True
    return p


def insert_picture(document: Document, png_path: Path, *, landscape: bool) -> None:
    usable_w = Mm(246) if landscape else Mm(160)
    usable_h = Mm(140) if landscape else Mm(180)
    with Image.open(png_path) as im:
        px_w, px_h = im.size
    aspect = px_w / px_h
    width = usable_w
    height = Emu(int(width) / aspect)
    if height > usable_h:
        height = usable_h
        width = Emu(int(height) * aspect)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_bidi(p, False)
    run = p.add_run()
    run.add_picture(str(png_path), width=width, height=height)


def add_figure_caption(document: Document, number: str, title: str) -> None:
    p = document.add_paragraph(style="Figure Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_bidi(p, True)
    add_mixed_text(p, f"איור {number} – {title}", size_pt=11, rtl_paragraph=True, base_bold=True)


def add_table_caption(document: Document, text: str) -> None:
    caption = text.strip().strip("*")
    p = document.add_paragraph(style="Table Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_bidi(p, True)
    add_mixed_text(p, caption, size_pt=11, rtl_paragraph=True, base_bold=True)


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    header, data = rows[0], rows[1:]
    table = document.add_table(rows=1 + len(data), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_bidi(table)
    prevent_table_split(table)
    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_bidi(p, True)
        p.paragraph_format.line_spacing = 1.0
        add_mixed_text(p, cell_text.strip(), size_pt=10, rtl_paragraph=True, base_bold=True)
        shade_cell(cell, "D9E2F3")
    for r_idx, row in enumerate(data, start=1):
        for c_idx in range(len(header)):
            value = row[c_idx].strip() if c_idx < len(row) else ""
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_bidi(p, True)
            p.paragraph_format.line_spacing = 1.0
            add_mixed_text(p, value, size_pt=10, rtl_paragraph=True)
            if r_idx % 2 == 0:
                shade_cell(cell, "F7F9FC")
    document.add_paragraph()


def parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_table_sep(line: str) -> bool:
    raw = line.strip().replace("|", "").replace(":", "").replace("-", "").replace(" ", "")
    return line.strip().startswith("|") and set(line) <= set("| -:")


def looks_hebrew(text: str) -> bool:
    return bool(re.search(r"[\u0590-\u05FF]", text))


def figure_key_from_text(text: str) -> str | None:
    for key in FIGURE_META:
        if key in text or f"{key}.mmd" in text:
            return key
    return None


def setup_headers_footers(document: Document) -> None:
    for i, section in enumerate(document.sections):
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.clear()
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_bidi(hp, True)
        add_mixed_text(hp, "The Bridge — דוח פרויקט גמר", size_pt=10, rtl_paragraph=True)
        for run in hp.runs:
            run.italic = True

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.clear()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_bidi(fp, False)
        run = fp.add_run("עמוד ")
        set_run_fonts(run, size_pt=10, cs_font=HE_FONT)
        set_run_rtl(run, True)
        add_page_number(fp)

    # Cover: different first page, blank header/footer
    first = document.sections[0]
    first.different_first_page_header_footer = True
    first.first_page_header.paragraphs[0].clear()
    first.first_page_footer.paragraphs[0].clear()


def add_update_note(document: Document) -> None:
    p = document.add_paragraph(style="Note")
    set_bidi(p, True)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_mixed_text(
        p,
        "יש לפתוח מסמך זה ב-Microsoft Word ולבחור Update Field / Update entire table "
        "עבור תוכן העניינים, רשימת האיורים ורשימת הטבלאות (לחיצה ימנית על השדה).",
        size_pt=10,
        rtl_paragraph=True,
        base_bold=True,
    )


def add_toc_block(document: Document, heading: str, instruction: str, placeholder: str) -> None:
    p = document.add_paragraph(style="TOCHeading")
    p.paragraph_format.page_break_before = True
    set_bidi(p, True)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_mixed_text(p, heading, size_pt=16, rtl_paragraph=True, base_bold=True)
    for run in p.runs:
        run.bold = True
    add_update_note(document)
    field_p = document.add_paragraph()
    set_bidi(field_p, True)
    field_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_field(field_p, instruction, placeholder)


def write_code_block(document: Document, lines: list[str]) -> None:
    for line in lines:
        p = document.add_paragraph(style="CodeBlock")
        set_bidi(p, False)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(line if line else " ")
        set_run_fonts(run, ascii_font=CODE_FONT, cs_font=CODE_FONT, size_pt=10)
        set_run_rtl(run, False)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        set_element_attr(shd, "fill", "F4F4F4")
        set_element_attr(shd, "val", "clear")
        pPr.append(shd)
    document.add_paragraph()


class ReportBuilder:
    def __init__(self, document: Document):
        self.document = document
        self.inserted_figures: set[str] = set()
        self.in_english_abstract = False
        self.chapter_started = False

    def switch_orientation(self, landscape: bool) -> None:
        current = self.document.sections[-1]
        already = current.page_width > current.page_height
        if already == landscape:
            return
        section = self.document.add_section(WD_SECTION.NEW_PAGE)
        configure_section(section, landscape=landscape)
        setup_headers_footers(self.document)

    def ensure_portrait(self) -> None:
        self.switch_orientation(False)

    def insert_figure(self, key: str) -> None:
        if key in self.inserted_figures:
            return
        number, title, landscape = FIGURE_META[key]
        png = RENDERED / f"{key}.png"
        if not png.exists():
            add_body_paragraph(self.document, f"[איור {number} חסר: {png.name}]")
            return
        self.switch_orientation(landscape)
        insert_picture(self.document, png, landscape=landscape)
        add_figure_caption(self.document, number, title)
        self.inserted_figures.add(key)
        self.switch_orientation(False)

    def add_heading(self, raw: str, level: int) -> None:
        text = raw.strip()
        if text in TOC_HEADINGS:
            return
        self.ensure_portrait()
        self.in_english_abstract = text == "Abstract"
        page_break = False
        if level == 1:
            if text in FRONT_HEADINGS:
                page_break = True
            elif re.match(r"^\d+\.", text) or text.startswith("נספח") or text == "מקורות":
                page_break = True
        heading_paragraph(self.document, text, level, page_break=page_break)

    def add_paragraph(self, text: str) -> None:
        text = text.strip()
        if not text or text == "---":
            return
        if SKIP_BODY_RE.search(text):
            return
        if text.startswith("![") and "](" in text:
            key = figure_key_from_text(text)
            if key:
                self.insert_figure(key)
            return
        key = figure_key_from_text(text)
        cleaned = re.sub(r"\s*מקור:\s*`[^`]+`\.?", "", text).strip()
        cleaned = re.sub(r"\s*`docs/diagrams/source/[^`]+`", "", cleaned).strip()
        cleaned = re.sub(r"\s*:\s*$", "", cleaned)
        rtl = not self.in_english_abstract
        style = "EnglishBody" if self.in_english_abstract else "Normal"
        if cleaned:
            add_body_paragraph(
                self.document,
                cleaned,
                style=style,
                rtl=rtl,
            )
        if key:
            self.insert_figure(key)


def build_cover(document: Document, lines: list[str]) -> None:
    # vertical space
    for _ in range(3):
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
    title = document.add_paragraph(style="CoverTitle")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_bidi(title, False)
    run = title.add_run("THE BRIDGE")
    set_run_fonts(run, ascii_font=EN_FONT, cs_font=EN_FONT, size_pt=28, bold=True)
    set_run_rtl(run, False)

    for line in lines:
        line = line.strip().strip("*")
        if not line or line == "THE BRIDGE" or line == "---":
            continue
        p = document.add_paragraph(style="CoverLine")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rtl = looks_hebrew(line)
        set_bidi(p, rtl)
        add_mixed_text(p, line, size_pt=14 if looks_hebrew(line) else 13, rtl_paragraph=rtl)


def parse_markdown(md: str) -> None:
    document = Document()
    configure_section(document.sections[0], landscape=False)
    setup_styles(document)
    set_update_fields_on_open(document)

    start = md.find("# THE BRIDGE")
    body = md[start:]
    raw_lines = body.splitlines()

    builder = ReportBuilder(document)

    # Cover: from THE BRIDGE until תקציר
    cover_lines: list[str] = []
    i = 0
    while i < len(raw_lines):
        if raw_lines[i].startswith("# תקציר"):
            break
        if raw_lines[i].startswith("# "):
            cover_lines.append(raw_lines[i][2:].strip())
        else:
            cover_lines.append(raw_lines[i])
        i += 1
    build_cover(document, cover_lines)

    # Front-matter TOCs will be inserted when those headings are seen
    toc_done = False

    while i < len(raw_lines):
        line = raw_lines[i]
        stripped = line.strip()

        if stripped.startswith("# ") and stripped[2:].strip() in TOC_HEADINGS:
            if not toc_done:
                add_toc_block(
                    document,
                    "תוכן עניינים",
                    r'TOC \o "1-3" \h \z \u',
                    "תוכן העניינים יעודכן בעת פתיחת הקובץ ב-Microsoft Word.",
                )
                add_toc_block(
                    document,
                    "רשימת איורים",
                    r'TOC \h \z \t "Figure Caption,1"',
                    "רשימת האיורים תעודכן ב-Microsoft Word (Update Field).",
                )
                add_toc_block(
                    document,
                    "רשימת טבלאות",
                    r'TOC \h \z \t "Table Caption,1"',
                    "רשימת הטבלאות תעודכן ב-Microsoft Word (Update Field).",
                )
                toc_done = True
            i += 1
            continue

        if stripped.startswith("```"):
            fence = []
            i += 1
            while i < len(raw_lines) and not raw_lines[i].strip().startswith("```"):
                fence.append(raw_lines[i])
                i += 1
            write_code_block(document, fence)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(raw_lines) and is_table_sep(raw_lines[i + 1]):
            rows = [parse_table_row(stripped)]
            i += 2
            while i < len(raw_lines) and raw_lines[i].strip().startswith("|"):
                rows.append(parse_table_row(raw_lines[i]))
                i += 1
            builder.ensure_portrait()
            add_markdown_table(document, rows)
            continue

        if stripped.startswith("### "):
            builder.add_heading(stripped[4:], 3)
            i += 1
            continue
        if stripped.startswith("## "):
            builder.add_heading(stripped[3:], 2)
            i += 1
            continue
        if stripped.startswith("# "):
            builder.add_heading(stripped[2:], 1)
            i += 1
            continue

        if re.match(r"^\*\*טבלה ", stripped):
            builder.ensure_portrait()
            add_table_caption(document, stripped)
            i += 1
            continue

        if re.match(r"^(\d+\.|-|\*)\s+", stripped):
            builder.ensure_portrait()
            p = document.add_paragraph(style="ListBulletHe")
            p.paragraph_format.right_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.75)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_bidi(p, True)
            add_mixed_text(p, stripped, size_pt=12, rtl_paragraph=True)
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if stripped:
            builder.ensure_portrait()
            builder.add_paragraph(stripped)
        i += 1

    # Ensure all figures exist even if a reference was missed
    for key in FIGURE_META:
        if key not in builder.inserted_figures:
            builder.insert_figure(key)

    builder.ensure_portrait()
    setup_headers_footers(document)
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT_PATH)


def verify(path: Path) -> None:
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        xml = zf.read("word/document.xml").decode("utf-8")
        media = [n for n in names if n.startswith("word/media/")]
        print("media files:", len(media), media)
        for token in ("**", "## ", "```", "!["):
            # raw markdown should not appear as visible text nodes
            if f">{token}" in xml or f">{token.rstrip()}":
                count = xml.count(token)
                print(f"token {token!r} xml occurrences:", count)
        print("bidi flags:", xml.count("<w:bidi"))
        print("TOC instr:", "TOC " in xml)
        print("PAGE field:", "PAGE" in xml)
        print("Heading1 count approx:", xml.count("Heading1") + xml.count("Heading 1"))

    doc = Document(str(path))
    print("paragraphs:", len(doc.paragraphs))
    print("tables:", len(doc.tables))
    print("sections:", len(doc.sections))
    heads = [p.text for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
    print("headings sample:", heads[:8], "... total", len(heads))
    captions_f = [p.text for p in doc.paragraphs if p.style and p.style.name == "Figure Caption"]
    captions_t = [p.text for p in doc.paragraphs if p.style and p.style.name == "Table Caption"]
    print("figure captions:", captions_f)
    print("table captions:", len(captions_t), captions_t[:3], "...")
    # leftover markdown headings
    leftover = [p.text for p in doc.paragraphs if p.text.strip().startswith("#")]
    print("leftover # paragraphs:", leftover)
    stars = [p.text for p in doc.paragraphs if "**" in p.text]
    print("paragraphs with **:", len(stars), stars[:3])


def main() -> None:
    md = MD_PATH.read_text(encoding="utf-8")
    parse_markdown(md)
    print("wrote", OUT_PATH, "size", OUT_PATH.stat().st_size)
    verify(OUT_PATH)


if __name__ == "__main__":
    main()
