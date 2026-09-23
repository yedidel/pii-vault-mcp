#!/usr/bin/env python3
"""Second-pass QA fixes on the READY submission DOCX only."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

DST = Path("/Users/yhd/Downloads/The_Bridge_Final_Report_READY_FOR_SUBMISSION.docx")
SEQ = Path("/Users/yhd/pii-vault-mcp/docs/diagrams/rendered/02-sequence.png")
HE = "David"
EN = "Arial"


def qn_w(tag: str) -> str:
    return qn(f"w:{tag}")


def set_val(el, name: str, value: str) -> None:
    el.set(qn(f"w:{name}"), value)


def delete_paragraph(paragraph: Paragraph) -> None:
    el = paragraph._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def insert_paragraph_before(paragraph: Paragraph, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    return new_para


def insert_paragraph_after(paragraph: Paragraph, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    return new_para


def set_run_fonts(run, *, size=12, bold=None, italic=None, cs=HE, ascii_font=EN, rtl=None):
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.size = Pt(size)
    rPr = run._r.get_or_add_rPr()
    for child in list(rPr):
        if child.tag == qn_w("rFonts"):
            rPr.remove(child)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:cs"), cs)
    rFonts.set(qn("w:eastAsia"), ascii_font)
    rPr.insert(0, rFonts)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:bidi"), "he-IL")
    rPr.append(lang)
    sz = OxmlElement("w:sz")
    set_val(sz, "val", str(int(size * 2)))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    set_val(szCs, "val", str(int(size * 2)))
    rPr.append(szCs)
    if rtl is not None:
        rtl_el = OxmlElement("w:rtl")
        set_val(rtl_el, "val", "1" if rtl else "0")
        rPr.append(rtl_el)


def set_para_bidi(p: Paragraph, rtl: bool) -> None:
    pPr = p._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag == qn_w("bidi"):
            pPr.remove(child)
    bidi = OxmlElement("w:bidi")
    set_val(bidi, "val", "1" if rtl else "0")
    pPr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 1.5


def rewrite(p: Paragraph, text: str, *, rtl=True, size=12, bold=False, style=None, italic=False) -> None:
    if style:
        p.style = style
    for child in list(p._p):
        if child.tag != qn_w("pPr"):
            p._p.remove(child)
    run = p.add_run(text)
    set_run_fonts(run, size=size, bold=bold, italic=italic, rtl=rtl)
    set_para_bidi(p, rtl)


def find_one(doc: Document, start: str) -> Paragraph:
    hits = [p for p in doc.paragraphs if p.text.strip().startswith(start)]
    if not hits:
        raise SystemExit(f"not found: {start}")
    return hits[0]


def ensure_caption_style(doc: Document) -> str:
    from docx.enum.style import WD_STYLE_TYPE

    try:
        st = doc.styles["Caption"]
    except KeyError:
        st = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
    st.font.size = Pt(11)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.name = EN
    rPr = st.element.get_or_add_rPr()
    rFonts = rPr.find(qn_w("rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), EN)
    rFonts.set(qn("w:hAnsi"), EN)
    rFonts.set(qn("w:cs"), HE)
    return "Caption"


def captionize(p: Paragraph) -> None:
    try:
        p.style = "Caption"
    except Exception:
        pass
    set_para_bidi(p, True)
    for r in p.runs:
        set_run_fonts(r, size=11, bold=True, rtl=True)


def insert_table_caption(table, text: str) -> None:
    new_p = OxmlElement("w:p")
    table._tbl.addprevious(new_p)
    para = Paragraph(new_p, table._parent)
    rewrite(para, text, size=11, bold=True, style="Caption")


def main() -> None:
    doc = Document(str(DST))
    ensure_caption_style(doc)

    leftover_toc = {
        "1. מבוא",
        "2. רקע תאורטי ומחקרי",
        "3. הגדרת הבעיה ומטרות הפרויקט",
        "4. דרישות המערכת",
        "5. חלופות ושיקולי תכנון",
        "6. ארכיטקטורת המערכת",
        "7. תכנון מפורט ומימוש",
        "8. מנוע ה-Workflow וניהול מצבים",
        "9. שכבת הסוכנים והמשך שיחה",
        "10. מודל הנתונים והתמדה",
        "10. מודל הנתונים ושמירתם",
        "11. ממשקי תקשורת",
        "12. אבטחה ופרטיות",
        "13. בדיקות והרצות",
        "14. תוצאות, מגבלות ודיון",
        "15. עבודה עתידית",
        "16. סיכום",
        "מקורות",
        "נספחים",
    }
    toc_h = find_one(doc, "תוכן עניינים")
    fig_h = find_one(doc, "רשימת איורים")
    between = True
    for p in list(doc.paragraphs):
        if p._p is toc_h._p:
            continue
        if p._p is fig_h._p:
            between = False
            continue
        if not between:
            continue
        t = p.text.strip()
        if t in leftover_toc:
            delete_paragraph(p)

    rewrite(
        find_one(doc, "A Secure Mediation"),
        "A Mediation and Process Orchestration Layer for AI Agents",
        rtl=False,
        size=14,
        bold=True,
    )
    rewrite(
        find_one(doc, "הפרויקט מדגים תרחיש מלא מקצה לקצה"),
        "הפרויקט מדגים תרחיש מלא מקצה לקצה של קבלת בקשה מסוכן, אימות הסוכן, המתנה לאישור או לתוצאה חיצונית, שמירת מצב, השלמת הפעולה והמשך שיחת הסוכן בלי לחשוף את ערכי המידע הרגיש לסוכן. התרומה המרכזית היא ארכיטקטורה מודולרית, בלתי תלויה בספק AI מסוים, המאפשרת לשלב סוכנים בתהליכים רגישים וארוכי טווח באב-טיפוס מקומי.",
    )
    rewrite(
        find_one(doc, "The project demonstrates an end-to-end"),
        "The project demonstrates an end-to-end flow that receives an agent request, verifies trust, waits for external approval, preserves durable state, completes the action, and resumes the agent conversation without exposing sensitive user data to the agent. Its main contribution is a modular, provider-independent architecture for asynchronous AI-agent integration. The current system is a local functional prototype.",
        rtl=False,
    )

    cap = find_one(doc, "איור 6.2:")
    img_p = insert_paragraph_before(cap)
    set_para_bidi(img_p, False)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_p.add_run()
    run.add_picture(str(SEQ), width=Inches(6.1))

    for start in ("איור 6.1:", "איור 6.2:", "איור 8.1:", "איור 10.1:"):
        captionize(find_one(doc, start))

    # Table captions (above tables)
    captions = [
        "טבלה 4.1: דרישות פונקציונליות",
        "טבלה 4.2: דרישות לא-פונקציונליות",
        "טבלה 6.1: רכיבי המערכת ואחריותם",
        "טבלה 7.1: חבילות Bridge ותפקידיהן",
        "טבלה 8.1: אירועי Workflow מרכזיים",
        "טבלה 11.1: ממשק gRPC",
        "טבלה 11.2: כלי MCP",
        "טבלה 11.3: נקודות HTTP, Webhook ו-SSE",
        "טבלה 13.1: תרחישי בדיקה צפויים (לא תוצאות הרצה חתומות)",
        "טבלה א.1: מפת רכיבים ומחלקות",
        "טבלה ב.1: מילון מונחים",
        "טבלה ג.1: משתני סביבה ותצורה",
    ]
    if len(doc.tables) != len(captions):
        raise SystemExit(f"table count {len(doc.tables)} != captions {len(captions)}")
    for tbl, text in zip(doc.tables, captions):
        insert_table_caption(tbl, text)

    http_h = find_one(doc, "11.4 HTTP")
    note = insert_paragraph_after(http_h)
    rewrite(
        note,
        "הטבלה שלהלן מאגדת נקודות HTTP של כמה תהליכים: Adapter במצב HTTP MCP, שרת FastAPI הנפרד במצב stdio (webhook/events), Claude Channel, ו-Install Directory. אין לראות את כל הנתיבים כשייכים לאותו תהליך או לאותו מודל אימות. HMAC ל-webhook מאומת במסלול HTTP MCP; במצב stdio נקודת POST /webhook/bridge מעבדת את הגוף בלי אימות חתימה.",
    )

    sdk = insert_paragraph_after(find_one(doc, "הודעה ללקוח המקומי באמצעות SSE"))
    rewrite(
        sdk,
        "קיים גם SdkProvider להמשך דרך לקוח ה-SDK; ה-Adapter עצמו אינו מייבא את חבילת bridge-sdk.",
    )

    replay = find_one(doc, "12.5.2")
    body = replay.text.split(":", 1)[-1].strip()
    rewrite(replay, "12.5.2 הגנה מפני שידור חוזר (Anti-Replay via Timestamp & Nonce)", style="Heading 3", size=12, bold=True)
    after = insert_paragraph_after(replay)
    rewrite(
        after,
        body
        + " מנגנון זה חל יחד עם אימות HMAC במסלול HTTP MCP. במסלול stdio, POST /webhook/bridge אינו מאמת חתימה ולכן גם בדיקות replay אינן חלות שם.",
    )

    # Split remaining long Heading 3 titles into heading + body
    for p in list(doc.paragraphs):
        if p.style.name != "Heading 3":
            continue
        t = p.text.strip()
        if ":" not in t or len(t) < 80:
            continue
        if t.startswith("12.5.2"):
            continue
        num, rest = t.split(":", 1)
        heading = num.strip()
        body = rest.strip()
        if body.startswith("(") and ")" in body[:80]:
            # keep short parenthetical in heading if it is a name
            close = body.find(")")
            heading = num.strip() + ": " + body[: close + 1]
            body = body[close + 1 :].lstrip(" :")
        rewrite(p, heading, style="Heading 3", size=12, bold=True)
        if body:
            bp = insert_paragraph_after(p)
            rewrite(bp, body)

    secrets = find_one(doc, "12.7.3")
    # after split, heading may be short
    if "load_adapter_env" in secrets.text or len(secrets.text) > 80:
        pass
    else:
        nxt = secrets._p.getnext()
        # body already following

    # Font pass for English cover/abstract
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith(
            (
                "A Mediation",
                "AI agents",
                "The Bridge is a mediation",
                "The implementation combines",
                "The project demonstrates",
            )
        ) or t == "THE BRIDGE":
            set_para_bidi(p, False)
            for r in p.runs:
                size = 18 if t == "THE BRIDGE" else (14 if t.startswith("A Mediation") else 12)
                set_run_fonts(r, size=size, bold=(t in {"THE BRIDGE"} or t.startswith("A Mediation")), ascii_font=EN, cs=EN, rtl=False)

    # Heading sizes
    for p in doc.paragraphs:
        if p.style.name == "Heading 1":
            if p.text.strip() == "Abstract":
                set_para_bidi(p, False)
                for r in p.runs:
                    set_run_fonts(r, size=16, bold=True, ascii_font=EN, cs=EN, rtl=False)
            else:
                set_para_bidi(p, True)
                for r in p.runs:
                    set_run_fonts(r, size=16, bold=True, rtl=True)
        elif p.style.name == "Heading 2":
            set_para_bidi(p, True)
            for r in p.runs:
                set_run_fonts(r, size=14, bold=True, rtl=True)
        elif p.style.name == "Heading 3":
            set_para_bidi(p, True)
            for r in p.runs:
                set_run_fonts(r, size=12, bold=True, rtl=True)

    doc.save(str(DST))
    print("updated", DST, DST.stat().st_size)


if __name__ == "__main__":
    main()
